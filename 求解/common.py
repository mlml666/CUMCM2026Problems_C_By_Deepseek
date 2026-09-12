# -*- coding: utf-8 -*-
"""CUMCM2026 C题 公共模块（Phase 2）
数据装载、LP 求解引擎、结算、结果文件写出。

口径（Phase 1 定死，见 求解/模型推导/00_符号与共同模型.md）：
- 时段 t=1..144，时段 t = 区间 ((t-1)/6 h, t/6 h]，即 [0:00,0:10] ... [23:50,24:00]
- 数据列标签 00:10 .. 23:50, 0:00+1 = 各时段末时刻（解读A）
- result 模板的 144 行/列按时间顺序一一对应填充（模板标签晚 10 分钟，属模板命名偏移，见验算报告）
- 储能：电池侧变量 a,b；E_t = E_{t-1} + a_t - b_t；母线平衡 g + η_d b + PΔt = LΔt + a/η_c + s
"""
import os
import datetime as dt
import numpy as np
import pandas as pd
from scipy.optimize import linprog
from scipy import sparse
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))


def _find_root(start):
    """向上查找包含 附件/ 的目录，使附录中的源码副本在任意布局下均可运行"""
    cur = start
    for _ in range(6):
        if os.path.isdir(os.path.join(cur, '附件')):
            return cur
        parent = os.path.dirname(cur)
        if parent == cur:
            break
        cur = parent
    return os.path.dirname(start)


def _pick_dir(*cands):
    for c in cands:
        if os.path.isdir(c):
            return c
    return cands[0]


ROOT = _find_root(HERE)
ATT = os.path.join(ROOT, '附件')
TPL = os.path.join(ATT, '附件5')
RES = _pick_dir(os.path.join(HERE, '结果'), os.path.join(ROOT, '求解', '结果'))
FIG = _pick_dir(os.path.join(HERE, '图片'), os.path.join(ROOT, '求解', '图片'))
os.makedirs(RES, exist_ok=True)
os.makedirs(FIG, exist_ok=True)

# ---------------- 常量（附录1 与 Phase 1 决策） ----------------
T = 144
DT = 1.0 / 6.0
ETA_C = 0.9
ETA_D = 0.9
E_MIN, E_MAX, E_CAP = 1200.0, 10800.0, 12000.0
P_B = 5000.0
ABAR = P_B * DT                 # 833.333 kWh/时段（电池侧）
E0 = 6000.0
EMERG_MULT = 5.0
DEV_DOWN, DEV_UP = 0.5, 1.5
SEED = 42
np.random.seed(SEED)

REPORT_START = dt.date(2025, 2, 1)
REPORT_END = dt.date(2025, 12, 31)
BLOCKS = ['0:00-4:00', '4:00-8:00', '8:00-12:00', '12:00-16:00', '16:00-20:00', '20:00-24:00']
SLOT_LABELS_TPL = None  # 由模板读取


# ---------------- 数据装载 ----------------
def load_att1():
    df = pd.read_excel(os.path.join(ATT, '附件1.xlsx'))
    df.columns = ['时间', '电价', '小区负载', '光伏发电预测功率']
    for c in df.columns[1:]:
        df[c] = pd.to_numeric(df[c], errors='coerce')
    return df


def _load_wide(fname, sheet=None):
    path = os.path.join(ATT, fname)
    df = pd.read_excel(path) if sheet is None else pd.read_excel(path, sheet_name=sheet)
    dates = pd.to_datetime(df.iloc[:, 0].astype(str).str[:10]).dt.date.values
    vals = df.iloc[:, 1:].astype(float).values
    return dates, vals


def load_att2():
    dl, load = _load_wide('附件2.xlsx', '小区负载')
    _, pv = _load_wide('附件2.xlsx', '光伏发电实际功率')
    return dl, load, pv


def load_att3():
    df = pd.read_excel(os.path.join(ATT, '附件3.xlsx'))
    df.columns = ['日期', '预报时刻'] + ['预报%d小时' % i for i in range(1, 25)]
    df['日期'] = pd.to_datetime(df['日期']).ffill().dt.date
    out = {}
    for _, r in df.iterrows():
        out[(r['日期'], str(r['预报时刻']))] = r.iloc[2:].astype(float).values
    return out


def load_att4():
    return _load_wide('附件4.xlsx')


def day_index(dates, d):
    idx = np.where(dates == d)[0]
    return int(idx[0]) if len(idx) else None


def fc_hourly_to_slots(hourly, u=0):
    """附件3 预报（发布时刻 u，未来第 k 小时 k=1..24）→ 当日 144 时段（分段常数，W3）

    关键口径：预报k小时 = 时钟 (u+k) 时；超出当日 24:00 的部分置 0。
    """
    v = np.zeros(T)
    h = np.asarray(hourly, dtype=float)
    for k in range(1, 25):
        hh = u + k
        if hh <= 24:
            v[(hh - 1) * 6: hh * 6] = h[k - 1]
    return v


def forecast_path_clock(fc, d, u, issues=(0, 6, 12, 18)):
    """按发布时刻 u 可得的"今日逐整点预报"（相邻发布时刻线性插值），返回长度 24（索引=时钟小时-1）"""
    v = np.full(24, np.nan)
    for t in issues:
        if t > u:
            continue
        h = fc.get((d, '%d:00' % t))
        if h is None:
            continue
        for k in range(1, 25):
            hh = t + k
            if hh <= 24:
                v[hh - 1] = h[k - 1]
    prev = max([t for t in issues if t <= u], default=None)
    nxt = min([t for t in issues if t > u], default=None)
    if prev is not None and nxt is not None:
        hp, hn = fc.get((d, '%d:00' % prev)), fc.get((d, '%d:00' % nxt))
        if hp is not None and hn is not None:
            w = (u - prev) / float(nxt - prev)
            for k in range(1, 25):
                hh = nxt + k
                idxp, idxn = hh - prev - 1, hh - nxt - 1
                if hh <= 24 and 0 <= idxp < 24:
                    v[hh - 1] = (1 - w) * hp[idxp] + (w * hn[idxn] if 0 <= idxn < 24 else 0.0)
    return np.where(np.isnan(v), 0.0, v)


def clock_path_to_slots(v):
    return np.repeat(np.asarray(v, dtype=float), 6)


def report_dates():
    d0, d1 = REPORT_START, REPORT_END
    return [d0 + dt.timedelta(days=i) for i in range((d1 - d0).days + 1)]


# ---------------- 预测 ----------------
def ma_forecast(hist, W):
    """hist: (n_days, 144) 历史实际；返回同时段 W 日均值（最后 W 天）"""
    return hist[-W:].mean(axis=0)


PREDICTORS = ['昨日持续', '3日均', '7日均', '14日均', '30日均', 'EWMA(0.7,14)', 'EWMA(0.5,7)', '同月均']


def pred(hist, hist_dates, name, target_date):
    """滚动预测器（只用 target_date 之前的历史），hist=(n,144)"""
    if name == '昨日持续':
        return hist[-1]
    if name == '3日均':
        return hist[-3:].mean(0)
    if name == '7日均':
        return hist[-7:].mean(0)
    if name == '14日均':
        return hist[-14:].mean(0)
    if name == '30日均':
        return hist[-30:].mean(0)
    if name == 'EWMA(0.7,14)':
        h = hist[-14:]; w = 0.7 ** np.arange(len(h) - 1, -1, -1)
        return (h * (w / w.sum())[:, None]).sum(0)
    if name == 'EWMA(0.5,7)':
        h = hist[-7:]; w = 0.5 ** np.arange(len(h) - 1, -1, -1)
        return (h * (w / w.sum())[:, None]).sum(0)
    if name == '同月均':
        m = [j for j, dd in enumerate(hist_dates) if dd.month == target_date.month]
        return hist[m].mean(0) if m else hist[-7:].mean(0)
    raise KeyError(name)


def best_load_predictor():
    """从问题2 的精度对比结果中取负载 MAE 最优的预测器（保证各问口径一致）"""
    p = os.path.join(RES, '问题2_预测器精度.csv')
    if os.path.exists(p):
        df = pd.read_csv(p)
        return str(df.loc[df['负载MAE(kW)'].idxmin(), '预测器'])
    return '7日均'


def best_pv_predictor():
    """从问题2 的联合标定结果中取最优组合的光伏预测器"""
    p = os.path.join(RES, '问题2_裕度标定.csv')
    if os.path.exists(p):
        df = pd.read_csv(p).sort_values('total_cost_jan')
        return str(df.iloc[0]['光伏预测器'])
    return '7日均'


# ---------------- LP 求解引擎 ----------------
def solve_plan(L, P, price, E_init=E0, E_term=E0, ghat=None, dev_pen=0.0):
    """求解计划 LP。

    L, P, price : 长度 n 的数组（可为带裕度的预测值）
    E_init      : 首时段前的储电量
    E_term      : 末时段后的储电量要求（None = 自由）
    ghat        : 承诺计划（长度 n），配 dev_pen 施加 dev_pen*price*|g-ghat|
    返回 dict(g, a, b, s, E, fun, status)
    """
    n = len(L)
    L = np.asarray(L, float); P = np.asarray(P, float); price = np.asarray(price, float)
    nv = 7 * n
    c = np.zeros(nv)
    c[0:n] = price
    use_dev = (ghat is not None) and dev_pen > 0
    if use_dev:
        ghat = np.asarray(ghat, float)
        c[5 * n:6 * n] = dev_pen * price
        c[6 * n:7 * n] = dev_pen * price

    rows, cols, vals, b_eq = [], [], [], []
    ridx = 0
    # (1) 母线平衡: g + η_d b - a/η_c - s = LΔt - PΔt
    for i in range(n):
        rows += [ridx, ridx, ridx, ridx]
        cols += [i, 2 * n + i, n + i, 3 * n + i]
        vals += [1.0, ETA_D, -1.0 / ETA_C, -1.0]
        b_eq.append(L[i] * DT - P[i] * DT)
        ridx += 1
    # (2) 储能动态: E_i - E_{i-1} - a_i + b_i = 0
    for i in range(n):
        rows.append(ridx); cols.append(4 * n + i); vals.append(1.0)
        if i == 0:
            rows.append(ridx); cols.append(n + i); vals.append(-1.0)
            rows.append(ridx); cols.append(2 * n + i); vals.append(1.0)
            b_eq.append(E_init)
        else:
            rows.append(ridx); cols.append(4 * n + i - 1); vals.append(-1.0)
            rows.append(ridx); cols.append(n + i); vals.append(-1.0)
            rows.append(ridx); cols.append(2 * n + i); vals.append(1.0)
            b_eq.append(0.0)
        ridx += 1
    # (3) 偏差: dp - dm - g = -ghat
    if use_dev:
        for i in range(n):
            rows += [ridx, ridx, ridx]
            cols += [5 * n + i, 6 * n + i, i]
            vals += [1.0, -1.0, -1.0]
            b_eq.append(-ghat[i])
            ridx += 1
    # (4) 终端电量
    if E_term is not None:
        rows.append(ridx); cols.append(4 * n + n - 1); vals.append(1.0)
        b_eq.append(float(E_term))
        ridx += 1

    A_eq = sparse.csr_matrix((vals, (rows, cols)), shape=(ridx, nv))
    bounds = ([(0, None)] * n + [(0, ABAR)] * n + [(0, ABAR)] * n + [(0, None)] * n
              + [(E_MIN, E_MAX)] * n + [(0, None)] * n + [(0, None)] * n)
    res = linprog(c, A_eq=A_eq, b_eq=np.array(b_eq), bounds=bounds, method='highs')
    if not res.success:
        raise RuntimeError('LP 不可行/失败: %s' % res.message)
    x = res.x
    return dict(g=x[0:n].copy(), a=x[n:2 * n].copy(), b=x[2 * n:3 * n].copy(),
                s=x[3 * n:4 * n].copy(), E=x[4 * n:5 * n].copy(),
                fun=float(res.fun), status=res.status)


# ---------------- 结算 ----------------
def emergency_amount(L, P, g, b):
    return np.maximum(0.0, np.asarray(L) * DT - np.asarray(g) - ETA_D * np.asarray(b) - np.asarray(P) * DT)


def planned_cost(g, price):
    return float(np.dot(price, g))


def emergency_cost(e, price):
    return float(EMERG_MULT * np.dot(price, e))


def merge_intervals(e, tol=1e-3):
    """把逐时段缺电量合并为连续时间段 [(起, 止, 电量)]，时间格式 H:MM"""
    out = []
    t = 0
    n = len(e)
    while t < n:
        if e[t] > tol:
            t0 = t
            while t < n and e[t] > tol:
                t += 1
            qty = float(e[t0:t].sum())

            def lab(m):
                h, mm = divmod(int(m), 60)
                return ('%d:%02d' % (h, mm)) if m < 24 * 60 else '24:00'
            out.append((lab(t0 * 10), lab(t * 10), qty))
        else:
            t += 1
    return out


# ---------------- 结果文件写出 ----------------
def _wb(name):
    return openpyxl.load_workbook(os.path.join(TPL, name))


def write_result1(path, g, a, b, E):
    wb = _wb('result1.xlsx')
    ws = wb['计划购电量']
    for k in range(T):
        ws.cell(2 + k, 2).value = float(g[k])
    ws2 = wb['充放电量']
    for blk in range(6):
        t0, t1 = blk * 24, (blk + 1) * 24
        ws2.cell(2 + blk, 2).value = float(a[t0:t1].sum() / ETA_C)   # 充电量（母线侧）
        ws2.cell(2 + blk, 3).value = float(ETA_D * b[t0:t1].sum())   # 放电量（母线侧）
    ws2.cell(2, 4).value = dt.time(0, 0); ws2.cell(2, 5).value = float(E0)
    ws2.cell(3, 4).value = '24:00'; ws2.cell(3, 5).value = float(E[-1])
    wb.save(path)


def write_wide_sheet(ws, dates, mat, qty, cost):
    """写 日期×144 的宽表（模板 A 列日期与表头已存在）"""
    for i, d in enumerate(dates):
        ws.cell(2 + i, 1).value = dt.datetime(d.year, d.month, d.day)
        for k in range(T):
            ws.cell(2 + i, 2 + k).value = float(mat[i, k])
        ws.cell(2 + i, 146).value = float(qty[i])
        ws.cell(2 + i, 147).value = float(cost[i])


def write_charge_discharge(ws, days, chg, dis, Esoc):
    """days: 日期列表; chg/dis: (nd,6) 母线侧电量; Esoc: (nd,145) 储电量轨迹"""
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row)
    r = 2
    for i, d in enumerate(days):
        for j in range(6):
            ws.cell(r + j, 1).value = dt.datetime(d.year, d.month, d.day) if j == 0 else None
            ws.cell(r + j, 2).value = BLOCKS[j]
            ws.cell(r + j, 3).value = float(chg[i, j])
            ws.cell(r + j, 4).value = float(dis[i, j])
        ws.cell(r, 5).value = dt.time(0, 0); ws.cell(r, 6).value = float(Esoc[i, 0])
        ws.cell(r + 1, 5).value = '24:00'; ws.cell(r + 1, 6).value = float(Esoc[i, 144])
        r += 6


def write_emergency(ws, em_rows):
    """em_rows: [(date, '起-止', qty)]，同日只在首行写日期"""
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row)
    r = 2
    last = None
    for d, span, qty in em_rows:
        ws.cell(r, 1).value = dt.datetime(d.year, d.month, d.day) if d != last else None
        ws.cell(r, 2).value = span
        ws.cell(r, 3).value = float(qty)
        last = d
        r += 1


def blocks_of(x):
    """144 段 → 6 个 4 小时分块之和"""
    return np.array([x[i * 24:(i + 1) * 24].sum() for i in range(6)])


def fmt(x, n=2):
    return ('%.' + str(n) + 'f') % x


def setup_matplotlib():
    """统一图表风格（仅 matplotlib，禁 seaborn/jet/rainbow）"""
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    plt.rcParams['font.sans-serif'] = ['Microsoft YaHei', 'SimHei', 'DejaVu Sans']
    plt.rcParams['axes.unicode_minus'] = False
    plt.rcParams['figure.dpi'] = 120
    plt.rcParams['savefig.bbox'] = 'tight'
    plt.rcParams['axes.grid'] = True
    plt.rcParams['grid.alpha'] = 0.3
    return plt


def hours_axis():
    return np.arange(1, T + 1) / 6.0
